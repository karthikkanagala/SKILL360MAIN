"""
Resume Parsing Module
Extracts text from PDF, DOCX, and TXT files with OCR fallback for scanned PDFs
"""
from PyPDF2 import PdfReader
from docx import Document
import io
import os


def extract_text_from_pdf(file):
    """
    Extract text from PDF file with OCR fallback for scanned documents
    
    Args:
        file: File-like object (BytesIO) or file path
        
    Returns:
        Extracted text string
    """
    # First try normal text extraction
    try:
        reader = PdfReader(file)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
        
        # If we got meaningful text, return it
        if text.strip() and len(text.strip()) > 20:
            return text
    except Exception as e:
        print(f"PyPDF2 extraction failed: {e}")
    
    # If normal extraction failed or got minimal text, try OCR
    try:
        return extract_text_with_ocr(file)
    except Exception as e:
        print(f"OCR extraction failed: {e}")
        return text if text else ""


def extract_text_with_ocr(file):
    """
    Extract text from scanned PDF using OCR
    
    Args:
        file: File-like object or file path
        
    Returns:
        Extracted text string
    """
    try:
        import pytesseract
        from pdf2image import convert_from_bytes, convert_from_path
        from PIL import Image
        
        # Configure Tesseract path for Windows
        import platform
        if platform.system() == 'Windows':
            tesseract_paths = [
                r'C:\Program Files\Tesseract-OCR\tesseract.exe',
                r'C:\Program Files (x86)\Tesseract-OCR\tesseract.exe',
                r'C:\Users\ramka\AppData\Local\Programs\Tesseract-OCR\tesseract.exe'
            ]
            for path in tesseract_paths:
                if os.path.exists(path):
                    pytesseract.pytesseract.tesseract_cmd = path
                    break
        
        # Check if tesseract is available
        try:
            pytesseract.get_tesseract_version()
        except Exception:
            print("Tesseract not installed - OCR unavailable")
            return ""
        
        # Configure Poppler path for Windows
        poppler_path = None
        if platform.system() == 'Windows':
            poppler_paths = [
                r'C:\Users\ramka\AppData\Local\poppler-25.07.0\Library\bin',
                r'C:\Program Files\poppler\Library\bin',
                r'C:\Program Files\poppler-24.08.0\Library\bin',
            ]
            for path in poppler_paths:
                if os.path.exists(path):
                    poppler_path = path
                    break
        
        # Convert PDF to images
        if hasattr(file, 'read'):
            # It's a file-like object (BytesIO)
            file.seek(0)  # Reset to beginning
            pdf_bytes = file.read()
            if poppler_path:
                images = convert_from_bytes(pdf_bytes, poppler_path=poppler_path)
            else:
                images = convert_from_bytes(pdf_bytes)
        else:
            # It's a file path
            if poppler_path:
                images = convert_from_path(file, poppler_path=poppler_path)
            else:
                images = convert_from_path(file)
        
        # Extract text from each page image using OCR
        text = ""
        for i, image in enumerate(images):
            try:
                page_text = pytesseract.image_to_string(image)
                text += page_text + "\n"
            except Exception as page_error:
                print(f"OCR failed for page {i+1}: {page_error}")
                continue
        
        return text
        
    except ImportError as e:
        print(f"OCR libraries not available: {e}")
        return ""
    except Exception as e:
        print(f"OCR processing error: {e}")
        return ""


def extract_text_from_docx(file):
    """
    Extract text from DOCX file
    
    Args:
        file: File path string or file-like object
        
    Returns:
        Extracted text string
    """
    doc = Document(file)
    paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
    
    # Also extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())
    
    return "\n".join(paragraphs)


def extract_text_from_txt(file):
    """
    Extract text from TXT file
    
    Args:
        file: File-like object with read() method
        
    Returns:
        Extracted text string
    """
    if hasattr(file, 'read'):
        content = file.read()
        if isinstance(content, bytes):
            return content.decode('utf-8')
        return content
    else:
        with open(file, 'r', encoding='utf-8') as f:
            return f.read()
